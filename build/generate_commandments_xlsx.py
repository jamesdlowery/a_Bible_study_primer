#!/usr/bin/env python3
"""Build the commandments, laws and traditions workbook for a release:
a_Bible_study_primer_commandments_laws_traditions_<version>.xlsx.

A matrix of every commandment, law, ordinance, teaching and tradition in
the KJV: the pre-Mosaic commands of Genesis, the Decalogue (with the three
numbering systems), the traditional 613 commandments of the Torah in
Maimonides' enumeration (248 positive, 365 negative) with KJV proof-texts,
other OT statutes and instituted ordinances, the imperatives of the
Prophets and wisdom books, the commands of Christ, the apostolic
instructions and church ordinances, and the traditions and customs the
text itself names (the tradition of the elders, Jewish custom, apostolic
tradition, OT "manner"). Each row carries its category, kind, the
classical Moral / Ceremonial / Civil division (Mosaic rows), notes, how the
New Testament treats it, whether it presupposes the Temple or the land,
the primary KJV reference and parallels, and the opening of the KJV verse.

Sources:
  * build/commandments/commandments.csv -- every row and every column
    except the verse text (edit that file, never the spreadsheet). IDs:
    E/GA/GN/GP pre-Mosaic, D Decalogue, P1-P248 / N1-N365 Mosaic, O other OT,
    W Prophets & Wisdom, C Christ, A apostolic, T traditions.
  * build/commandments/kjv_verses.json.gz -- the KJV text (public domain),
    book -> chapter -> verse, in KJV chapter/verse numbering. Every
    reference in the CSV, primary and parallel, is resolved against it on
    every build, and the build fails if one does not resolve, so a typo in
    a reference can never ship; the opening of each primary reference is
    quoted into the workbook from it.

The workbook has a Summary (live COUNTIF formulas and a sheet index),
Read Me, Legend, the Master Matrix, and seven view tabs generated from the
Master Matrix with the same IDs and columns. Usage:
    python3 build/generate_commandments_xlsx.py <version> <repo_root> [out_path]
"""
import csv
import datetime
import gzip
import json
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

CSV_PATH = os.path.join("build", "commandments", "commandments.csv")
KJV_PATH = os.path.join("build", "commandments", "kjv_verses.json.gz")

ABBR = {
    "Genesis": "Gen", "Exodus": "Exod", "Leviticus": "Lev", "Numbers": "Num", "Deuteronomy": "Deut",
    "Joshua": "Josh", "Judges": "Judg", "Ruth": "Ruth", "1 Samuel": "1 Sam", "2 Samuel": "2 Sam",
    "1 Kings": "1 Kgs", "2 Kings": "2 Kgs", "1 Chronicles": "1 Chr", "2 Chronicles": "2 Chr", "Ezra": "Ezra",
    "Nehemiah": "Neh", "Esther": "Esth", "Job": "Job", "Psalms": "Ps", "Proverbs": "Prov", "Ecclesiastes": "Eccl",
    "Song of Solomon": "Song", "Isaiah": "Isa", "Jeremiah": "Jer", "Lamentations": "Lam", "Ezekiel": "Ezek",
    "Daniel": "Dan", "Hosea": "Hos", "Joel": "Joel", "Amos": "Amos", "Obadiah": "Obad", "Jonah": "Jonah",
    "Micah": "Mic", "Nahum": "Nah", "Habakkuk": "Hab", "Zephaniah": "Zeph", "Haggai": "Hag", "Zechariah": "Zech",
    "Malachi": "Mal", "Matthew": "Matt", "Mark": "Mark", "Luke": "Luke", "John": "John", "Acts": "Acts",
    "Romans": "Rom", "1 Corinthians": "1 Cor", "2 Corinthians": "2 Cor", "Galatians": "Gal", "Ephesians": "Eph",
    "Philippians": "Phil", "Colossians": "Col", "1 Thessalonians": "1 Thess", "2 Thessalonians": "2 Thess",
    "1 Timothy": "1 Tim", "2 Timothy": "2 Tim", "Titus": "Titus", "Philemon": "Phlm", "Hebrews": "Heb",
    "James": "Jas", "1 Peter": "1 Pet", "2 Peter": "2 Pet", "1 John": "1 John", "2 John": "2 John",
    "3 John": "3 John", "Jude": "Jude", "Revelation": "Rev",
}
NT_BOOKS = {b for b in ABBR if list(ABBR).index(b) >= list(ABBR).index("Matthew")}

DIV = {
    "Ten Commandments": "Moral", "Faith & worship of God": "Moral", "Idolatry & false worship": "Moral",
    "Neighbour, kindness & speech": "Moral", "Family, marriage & sexual": "Moral", "Oaths & vows": "Moral",
    "Sanctuary & priesthood": "Ceremonial", "Offerings & sacrifices": "Ceremonial", "Purity & uncleanness": "Ceremonial",
    "Dietary": "Ceremonial — dietary", "Sabbath & feasts": "Ceremonial — appointed times",
    "Tithes, firstfruits & dues": "Ceremonial", "Vows & Nazirite": "Ceremonial", "Torah study & signs": "Ceremonial",
    "Government, courts & justice": "Civil / judicial", "Criminal law & penalties": "Civil / judicial",
    "Civil law, property & commerce": "Civil / judicial", "Servants & slaves": "Civil / judicial",
    "Kings, war & nations": "Civil / judicial", "Agriculture & land": "Civil / judicial", "Poor, stranger & charity": "Civil / judicial",
}
DIV_FILL = {"Moral": "FFD966", "Ceremonial": "9DC3E6", "Ceremonial — dietary": "F8CBAD",
            "Ceremonial — appointed times": "C6E0B4", "Civil / judicial": "A9D18E"}
SEC_FILL = {
    "Pre-Mosaic (Genesis)": "FDEBD0",
    "Decalogue (Ten Commandments)": "FFF2CC",
    "Mosaic Law (Torah) — positive commands": "E2EFDA",
    "Mosaic Law (Torah) — negative commands": "FCE4E4",
    "Other OT statutes, covenant signs & instituted ordinances": "EDEDED",
    "Prophets & Wisdom — imperatives and summaries of the law": "FBE5D6",
    "New Testament — commands and teachings of Christ": "DDEBF7",
    "New Testament — apostolic instructions & church ordinances": "E4DFEC",
    "Traditions & customs": "F2F2F2",
}
KIND_FONT = {
    "Positive command": "1E6B2E", "Command of Christ": "1E6B2E", "Apostolic instruction": "1E6B2E",
    "Apostolic decree": "1E6B2E", "Imperative": "1E6B2E", "Imperative / summary": "1E6B2E",
    "Prohibition": "9C0006", "Prohibition / Judgment": "9C0006",
    "Judgment": "7F3F00", "Catalogue of sins": "7F3F00",
    "Teaching / principle": "1F4E79", "Principle": "1F4E79", "Summary / principle": "1F4E79",
    "Promise / covenant": "1F4E79", "Catalogue of righteousness": "1F4E79",
    "Ordinance": "2F5496", "Ordinance (instituted)": "2F5496", "Statute": "2F5496", "Covenant sign": "2F5496",
    "Church ordinance / practice": "2F5496", "Implied ordinance": "2F5496", "Permission": "2F5496", "Prologue": "2F5496",
    "Tradition (human / rabbinic)": "7030A0", "Apostolic tradition": "7030A0", "Custom / practice": "7030A0", "Custom": "7030A0",
}

HEADERS = ["ID", "Testament", "Section", "Category", "Kind", "Classical division",
           "Commandment / law / tradition", "Notes",
           "NT relation (OT rows: how the NT treats it · NT rows: relation to Mosaic law)",
           "Status (Mosaic rows)", "Primary reference (KJV)", "Parallel references", "Addressed to",
           "KJV text (primary reference, opening)"]
WIDTHS = [7, 13, 30, 26, 20, 18, 58, 34, 44, 22, 20, 30, 20, 62]
C = {h: i for i, h in enumerate(HEADERS)}
HR = 5  # header row on every table sheet (rows 1-3 are the title block)

FONT = "Arial"
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
HDR_FILL = PatternFill("solid", fgColor="1F3864")
BODY = Font(name=FONT, size=9)
BOLD = Font(name=FONT, size=9, bold=True)
_thin = Side(style="thin", color="D9D9D9")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
WRAP = Alignment(wrap_text=True, vertical="top")
TITLE = Font(name=FONT, bold=True, size=14, color="1F3864")
SUB = Font(name=FONT, italic=True, size=9, color="595959")

VIEWS = [
    ("Ten Commandments", "Exodus 20 / Deuteronomy 5, with the three numbering systems in the Notes column",
     lambda r: r["ID"].startswith("D")),
    ("613 Mitzvot", "The 248 positive (P) and 365 negative (N) commandments, Maimonides' enumeration, with KJV proof-texts",
     lambda r: r["ID"][0] in "PN" and r["Section"].startswith("Mosaic")),
    ("Pre-Mosaic & Other OT", "Genesis commands and customs (E/GA/GN/GP), other OT statutes and instituted ordinances (O)",
     lambda r: r["Section"].startswith(("Pre-Mosaic", "Other OT"))),
    ("Prophets & Wisdom", "Imperatives and summaries of the law in the Prophets, Psalms and wisdom books (W)",
     lambda r: r["ID"].startswith("W")),
    ("Christ's Commands", "Direct commands and governing teachings of Jesus; John the Baptist; Revelation 2–3 (C)",
     lambda r: r["ID"].startswith("C")),
    ("Apostolic", "Acts 15 decree, epistolary instructions and church ordinances (A)",
     lambda r: r["ID"].startswith("A") and r["Section"].startswith("New Testament")),
    ("Traditions & Customs", "Tradition of the elders, Jewish custom, apostolic tradition, OT 'manner' (T)",
     lambda r: r["ID"].startswith("T")),
]


# ---------------------------------------------------------------- references
_REF = re.compile(r"^((?:[1-3] )?[A-Za-z ]+?) (\d+):(\d+(?:-\d+)?(?:,\d+(?:-\d+)?)*)$")
_CHAP = re.compile(r"^((?:[1-3] )?[A-Za-z ]+?) (\d+)$")


def load_kjv(repo_root):
    with gzip.open(os.path.join(repo_root, KJV_PATH), "rt", encoding="utf-8") as fh:
        return json.load(fh)


def parse_refs(refstr, bible):
    """'Exodus 20:3; 22:1-4; Deuteronomy 5:7' -> [(book, chap, verse), ...].
    A piece without a book name continues the previous book; a parenthesis
    is a note and is ignored; 'Numbers 6' means that chapter (verse 1)."""
    refstr = re.sub(r"\s*\([^)]*\)", "", refstr)
    verses, last = [], None
    for piece in [p.strip() for p in refstr.split(";") if p.strip()]:
        m = _REF.match(piece)
        if not m and last and re.match(r"^\d+:\d", piece):
            m = _REF.match(f"{last} {piece}")
        if not m:
            mc = _CHAP.match(piece)
            if mc and mc.group(1) in bible and mc.group(2) in bible[mc.group(1)]:
                verses.append((mc.group(1), int(mc.group(2)), 1))
                last = mc.group(1)
                continue
            raise ValueError(f"unparsable reference {piece!r}")
        book = m.group(1)
        if book not in bible:
            raise ValueError(f"unknown book in {piece!r}")
        last = book
        chap = int(m.group(2))
        for part in m.group(3).split(","):
            a, b = (map(int, part.split("-")) if "-" in part else (int(part), int(part)))
            verses += [(book, chap, v) for v in range(a, b + 1)]
    return verses


def verse(bible, b, c, v):
    return bible.get(b, {}).get(str(c), {}).get(str(v))


def verify(rows, bible):
    errors = []
    for r in rows:
        for label in ("Primary reference", "Parallel references"):
            if not r[label]:
                continue
            try:
                for (b, c, v) in parse_refs(r[label], bible):
                    if verse(bible, b, c, v) is None:
                        errors.append(f"{r['ID']} {label}: {b} {c}:{v} is not in the KJV")
            except ValueError as e:
                errors.append(f"{r['ID']} {label}: {e}")
    return errors


def excerpt(bible, ref):
    parts = [verse(bible, b, c, v) for (b, c, v) in parse_refs(ref, bible)[:2]]
    text = " ".join(p for p in parts if p).replace("[", "").replace("]", "")
    return text[:317].rsplit(" ", 1)[0] + "..." if len(text) > 320 else text


def abbr(refs):
    out = []
    for p in [x.strip() for x in refs.split(";") if x.strip()]:
        m = re.match(r"^((?:[1-3] )?[A-Za-z ]+?) (\d+(?::.*)?)$", p)
        out.append(f"{ABBR.get(m.group(1), m.group(1))} {m.group(2)}" if m else p)
    return "; ".join(out)


def book_of(ref):
    return re.match(r"^((?:[1-3] )?[A-Za-z ]+?) \d", ref).group(1)


# ---------------------------------------------------------------- workbook
def header_block(ws, title, scope, nrows, version):
    ws["A1"] = title
    ws["A1"].font = TITLE
    ws["A2"] = scope
    ws["A2"].font = SUB
    ws["A3"] = (f"{nrows} rows · every reference verified against the KJV text at build time · "
                f"built with a Bible study primer {version}, {datetime.date.today().isoformat()}")
    ws["A3"].font = SUB
    for rr in (1, 2, 3):
        ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=8)


def write_table(ws, records):
    for i, h in enumerate(HEADERS, 1):
        c = ws.cell(HR, i, h)
        c.font, c.fill, c.border = HDR_FONT, HDR_FILL, BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[HR].height = 40
    for rec in records:
        ws.append(rec)
        row = ws[ws.max_row]
        fill = PatternFill("solid", fgColor=SEC_FILL.get(rec[C["Section"]], "FFFFFF"))
        kcol = KIND_FONT.get(rec[C["Kind"]], "000000")
        for c in row:
            c.font, c.alignment, c.border, c.fill = BODY, WRAP, BORDER, fill
        for h in ("ID", "Kind"):
            row[C[h]].font = Font(name=FONT, size=9, bold=True, color=kcol)
        row[C["Commandment / law / tradition"]].font = Font(name=FONT, size=9, color=kcol)
        div = rec[C["Classical division"]]
        if div in DIV_FILL:
            row[C["Classical division"]].fill = PatternFill("solid", fgColor=DIV_FILL[div])
            row[C["Classical division"]].font = BOLD
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(HR + 1, 2)
    ws.auto_filter.ref = f"A{HR}:{get_column_letter(len(HEADERS))}{ws.max_row}"


def hdr(ws, row, col, text):
    c = ws.cell(row, col, text)
    c.font, c.fill = HDR_FONT, HDR_FILL


def build(version, repo_root, out_path):
    with open(os.path.join(repo_root, CSV_PATH), encoding="utf-8", newline="") as fh:
        rows = list(csv.DictReader(fh))
    bible = load_kjv(repo_root)
    errors = verify(rows, bible)
    if errors:
        sys.exit("commandments.csv has references that do not resolve against the KJV:\n  " + "\n  ".join(errors))
    ids = [r["ID"] for r in rows]
    if len(ids) != len(set(ids)):
        sys.exit("commandments.csv has duplicate IDs")

    recs = []
    for r in rows:
        mosaic = r["Section"].startswith(("Decalogue", "Mosaic"))
        recs.append([
            r["ID"], "New Testament" if book_of(r["Primary reference"]) in NT_BOOKS else "Old Testament",
            r["Section"], r["Category"], r["Kind"], DIV.get(r["Category"], "—") if mosaic else "—",
            r["Description"], r["Notes"], r["NT relation"], r["Status"] or "—",
            r["Primary reference"], abbr(r["Parallel references"]), r["Addressed to"],
            excerpt(bible, r["Primary reference"]),
        ])
    by_id = {rec[0]: r for rec, r in zip(recs, rows)}

    wb = Workbook()
    wb.remove(wb.active)

    mm = wb.create_sheet("Master Matrix")
    header_block(mm, "Master Matrix — every commandment, law, ordinance and tradition",
                 "All sections in one filterable grid. The view tabs are generated from this sheet and carry the same IDs.",
                 len(recs), version)
    write_table(mm, recs)
    N = mm.max_row
    MM = "'Master Matrix'"

    view_counts = []
    for name, scope, pred in VIEWS:
        sub = [rec for rec in recs if pred(by_id[rec[0]])]
        ws = wb.create_sheet(name)
        header_block(ws, name, scope, len(sub), version)
        write_table(ws, sub)
        view_counts.append((name, len(sub)))
    if sum(n for _, n in view_counts) != len(recs):
        sys.exit(f"view tabs do not partition the rows: {view_counts} vs {len(recs)}")

    # ---- Summary: live formulas against the Master Matrix
    ss = wb.create_sheet("Summary", 0)
    ss["A1"] = "KJV Commandments, Laws & Traditions — Summary"
    ss["A1"].font = TITLE
    ss["A2"] = "All counts are COUNTIF(S) formulas over the Master Matrix; they update if rows are added or recategorised."
    ss["A2"].font = SUB
    ss["A3"] = f"Built with a Bible study primer {version}, {datetime.date.today().isoformat()}"
    ss["A3"].font = SUB
    sections, kinds, cats, statuses = [], [], [], []
    for rec in recs:
        if rec[C["Section"]] not in sections:
            sections.append(rec[C["Section"]])
        if rec[C["Kind"]] not in kinds:
            kinds.append(rec[C["Kind"]])
        if rec[C["Section"]].startswith("Mosaic"):
            if rec[C["Category"]] not in cats:
                cats.append(rec[C["Category"]])
            if rec[C["Status (Mosaic rows)"]] not in statuses:
                statuses.append(rec[C["Status (Mosaic rows)"]])
    rng = lambda col: f"{MM}!${col}${HR + 1}:${col}${N}"  # noqa: E731

    hdr(ss, 5, 1, "Section")
    hdr(ss, 5, 2, "Rows")
    r0 = 6
    for i, s in enumerate(sections):
        ss.cell(r0 + i, 1, s).font = BODY
        ss.cell(r0 + i, 2, f"=COUNTIF({rng('C')},A{r0 + i})").font = BODY
        for col in (1, 2):
            ss.cell(r0 + i, col).fill = PatternFill("solid", fgColor=SEC_FILL[s])
    tot = r0 + len(sections)
    ss.cell(tot, 1, "Total").font = BOLD
    ss.cell(tot, 2, f"=SUM(B{r0}:B{tot - 1})").font = BOLD
    k0 = tot + 2
    hdr(ss, k0, 1, "Kind")
    hdr(ss, k0, 2, "Rows")
    for i, k in enumerate(kinds):
        ss.cell(k0 + 1 + i, 1, k).font = Font(name=FONT, size=9, bold=True, color=KIND_FONT.get(k, "000000"))
        ss.cell(k0 + 1 + i, 2, f"=COUNTIF({rng('E')},A{k0 + 1 + i})").font = BODY

    m0 = 5
    for col, t in zip(range(4, 9), ["Mosaic Law (613) by category", "Positive", "Prohibition", "Total", "Classical division"]):
        hdr(ss, m0, col, t)
    for i, cat in enumerate(cats):
        rr = m0 + 1 + i
        ss.cell(rr, 4, cat).font = BODY
        ss.cell(rr, 5, f'=COUNTIFS({rng("D")},D{rr},{rng("E")},"Positive command",{rng("C")},"Mosaic*")').font = BODY
        ss.cell(rr, 6, f'=COUNTIFS({rng("D")},D{rr},{rng("E")},"Prohibition",{rng("C")},"Mosaic*")').font = BODY
        ss.cell(rr, 7, f"=E{rr}+F{rr}").font = BODY
        ss.cell(rr, 8, DIV[cat]).font = BODY
        ss.cell(rr, 8).fill = PatternFill("solid", fgColor=DIV_FILL[DIV[cat]])
    rt = m0 + 1 + len(cats)
    ss.cell(rt, 4, "Total").font = BOLD
    for col, L in ((5, "E"), (6, "F"), (7, "G")):
        ss.cell(rt, col, f"=SUM({L}{m0 + 1}:{L}{rt - 1})").font = BOLD
    d0 = rt + 2
    hdr(ss, d0, 4, "Decalogue + Mosaic Law by classical division")
    hdr(ss, d0, 5, "Rows")
    for i, d in enumerate(DIV_FILL):
        ss.cell(d0 + 1 + i, 4, d).font = BODY
        ss.cell(d0 + 1 + i, 4).fill = PatternFill("solid", fgColor=DIV_FILL[d])
        ss.cell(d0 + 1 + i, 5, f'=COUNTIFS({rng("F")},D{d0 + 1 + i},{rng("C")},"Mosaic*")'
                                f'+COUNTIFS({rng("F")},D{d0 + 1 + i},{rng("C")},"Decalogue*")').font = BODY
    s0 = d0 + len(DIV_FILL) + 2
    hdr(ss, s0, 4, "Mosaic Law by status (practicability today)")
    hdr(ss, s0, 5, "Rows")
    for i, st in enumerate(statuses):
        ss.cell(s0 + 1 + i, 4, st).font = BODY
        ss.cell(s0 + 1 + i, 5, f'=COUNTIFS({rng("J")},D{s0 + 1 + i},{rng("C")},"Mosaic*")').font = BODY
    x0 = max(k0 + len(kinds), s0 + len(statuses)) + 3
    hdr(ss, x0, 1, "Sheets in this workbook")
    hdr(ss, x0, 2, "Rows")
    index_rows = [("Summary", "Counts (live formulas) and this index"), ("Read Me", "Method, sections, columns, caveats"),
                  ("Legend", "Colour coding"), ("Master Matrix", "All rows, all columns, autofilter")]
    index_rows += [(n, "View of the Master Matrix") for n, _ in view_counts]
    for i, (name, note) in enumerate(index_rows):
        ss.cell(x0 + 1 + i, 1, f"{name} — {note}").font = BODY
        if name == "Master Matrix":
            ss.cell(x0 + 1 + i, 2, f"=COUNTA({rng('A')})").font = BODY
        for n, cnt in view_counts:
            if n == name:
                ss.cell(x0 + 1 + i, 2, cnt).font = BODY
    for col, w in zip("ABCDEFGH", [58, 8, 3, 40, 10, 12, 8, 26]):
        ss.column_dimensions[col].width = w

    # ---- Read Me
    rm = wb.create_sheet("Read Me", 1)
    notes = [
        ("KJV Commandments, Laws & Traditions — Read Me", "T"),
        (f"Built with a Bible study primer {version}. {len(recs)} rows. Every reference, primary and parallel, is resolved "
         f"against the full KJV text (66 books, 31,102 verses) on every build, and the build fails if one does not resolve; "
         f"the 'KJV text' column quotes the opening of the primary reference. Chapter/verse numbers follow the KJV, not the "
         f"Hebrew Bible (they differ in Exodus 20 & 22, Leviticus 5–6, Numbers 30, Deuteronomy 12–13 & 23).", ""),
        ("", ""),
        ("Source", "H"),
        ("The rows live in build/commandments/commandments.csv in the book's repository; edit that file, not this spreadsheet. "
         "The verse text comes from build/commandments/kjv_verses.json.gz.", ""),
        ("", ""),
        ("Sheets", "H"),
        ("Summary — live counts by section, kind, category, classical division and status, plus a sheet index.  Legend — the "
         "three colour layers.  Master Matrix — every row with every column and autofilter.  View tabs (Ten Commandments, 613 "
         "Mitzvot, Pre-Mosaic & Other OT, Prophets & Wisdom, Christ's Commands, Apostolic, Traditions & Customs) are generated "
         "from the Master Matrix and carry the same IDs, so each row has one source of truth.", ""),
        ("", ""),
        ("Sections (ID prefix)", "H"),
        ("E/GA/GN/GP Pre-Mosaic — Edenic, Adamic, Noahic, Patriarchal commands and customs in Genesis.", ""),
        ("D Decalogue — the Ten Commandments; the three numbering systems (Jewish; Catholic/Lutheran; Reformed/Orthodox/Anglican) "
         "are in the Notes column.", ""),
        ("P1–P248 / N1–N365 Mosaic Law — the traditional 613 commandments as enumerated by Maimonides (Sefer HaMitzvot), each "
         "with its primary KJV proof-text. A few are rabbinic derivations rather than explicit verses; the description says so "
         "where the verse is broader than the command.", ""),
        ("O Other OT — covenant signs, national statutes and later instituted ordinances outside the 613 (Purim, post-exilic "
         "fasts, Nehemiah's covenant).", ""),
        ("W Prophets & Wisdom — imperatives and summaries of the law in the Prophets, Psalms, Proverbs and Ecclesiastes.", ""),
        ("C Christ — direct commands and governing teachings of Jesus (and John the Baptist), and the risen Christ's commands "
         "to the seven churches.", ""),
        ("A Apostolic — the Jerusalem decree (Acts 15), instructions in the epistles, and church ordinances/practices recorded "
         "in Acts.", ""),
        ("T Traditions — things the text itself calls tradition, custom or 'manner': the tradition of the elders Jesus "
         "confronted, Second-Temple Jewish custom, apostolic tradition (paradosis), and OT customs.", ""),
        ("", ""),
        ("Columns", "H"),
        ("Kind — Positive command · Prohibition · Ordinance · Statute · Judgment · Covenant sign · Command of Christ · "
         "Teaching/principle · Imperative · Principle · Apostolic decree · Apostolic instruction · Church ordinance/practice · "
         "Tradition · Apostolic tradition · Custom.", ""),
        ("Classical division (Decalogue and Mosaic rows only) — Moral / Ceremonial / Civil (judicial), the threefold scheme of "
         "Aquinas, Calvin and the Westminster Confession ch. 19. Dietary laws and appointed times (Sabbath & feasts) are tagged "
         "as sub-bands of Ceremonial because they are the ceremonial laws most debated between denominations today. The scheme "
         "is a later interpretive framework, not a distinction the Torah draws; borderline cases are disputed.", ""),
        ("Notes — numbering systems, derivation flags, and 'Cited in NT' cross-references for Mosaic rows.", ""),
        ("NT relation — for OT rows, how the New Testament treats the law (Reaffirmed · Intensified · Reinterpreted · "
         "Fulfilled in Christ · Set aside · Not carried forward · Retained for Gentiles · Cited/echoed); for NT rows, how the "
         "command relates to the Mosaic law (Reaffirms · Intensifies · Extends · Restricts · Supersedes · Sets aside · Retains · "
         "New). Category-level defaults are overridden row by row where the NT addresses a specific law. These are descriptive "
         "of what the texts say; the theological weight given to them differs by tradition.", ""),
        ("Status (Mosaic rows) — Practicable anywhere · Temple-dependent · Land-dependent · Nation/land-dependent · Requires "
         "Israelite courts. Indicates whether the command presupposes the Temple, the land of Israel, or Israelite civil "
         "institutions.", ""),
        ("Parallel references use SBL-style abbreviations (Deut, Matt, 1 Cor…); the primary reference is spelled out so it "
         "sorts and filters cleanly.", ""),
        ("", ""),
        ("Limits", "H"),
        ("The 613 enumeration is one tradition (Nahmanides and Sefer HaChinuch differ slightly). There is no fixed canonical "
         "count of NT or prophetic commands; sections C, A and W aim for every distinct imperative without listing every "
         "repetition. The NT-relation and status tags are summaries, not verdicts — they are where denominational differences "
         "live, and the book can quote them as the texts' own data.", ""),
    ]
    for i, (t, h) in enumerate(notes, 1):
        c = rm.cell(i, 1, t)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = TITLE if h == "T" else (Font(name=FONT, bold=True, size=11) if h == "H" else Font(name=FONT, size=10))
    rm.column_dimensions["A"].width = 125

    # ---- Legend
    lg = wb.create_sheet("Legend", 2)
    lg["A1"] = "Colour legend"
    lg["A1"].font = TITLE
    lg["A2"] = ("Three independent colour layers are used on the Master Matrix and every view tab, so that each row can be "
                "read at a glance.")
    lg["A2"].font = SUB
    r = [4]

    def head(text):
        hdr(lg, r[0], 1, text)
        hdr(lg, r[0], 2, "Meaning")
        r[0] += 1

    head("1. Row shading = SECTION (where in the Bible / which body of law)")
    for s, note in {
        "Pre-Mosaic (Genesis)": "Commands and customs before Sinai — Eden, Adam, Noah, the patriarchs",
        "Decalogue (Ten Commandments)": "Exodus 20 / Deuteronomy 5",
        "Mosaic Law (Torah) — positive commands": "The 248 'thou shalt' commands of the 613",
        "Mosaic Law (Torah) — negative commands": "The 365 'thou shalt not' commands of the 613",
        "Other OT statutes, covenant signs & instituted ordinances": "Outside the 613: covenant signs, national statutes, Purim, post-exilic fasts",
        "Prophets & Wisdom — imperatives and summaries of the law": "Prophets, Psalms, Proverbs, Ecclesiastes",
        "New Testament — commands and teachings of Christ": "Gospels, Acts 1, Revelation 2–3",
        "New Testament — apostolic instructions & church ordinances": "Acts 15 decree, the epistles, church practice in Acts",
        "Traditions & customs": "Tradition of the elders, Jewish custom, apostolic tradition, OT 'manner'",
    }.items():
        c = lg.cell(r[0], 1, s)
        c.fill, c.font, c.border = PatternFill("solid", fgColor=SEC_FILL[s]), BODY, BORDER
        lg.cell(r[0], 2, note).font = BODY
        r[0] += 1
    r[0] += 1
    head("2. Font colour = KIND (what sort of rule it is) — on the ID, Kind and Description columns")
    for color, label, text in [
        ("1E6B2E", "Dark green", "DO: positive commands, commands of Christ, apostolic instructions and decrees, prophetic/wisdom imperatives"),
        ("9C0006", "Dark red", "DO NOT: prohibitions"),
        ("2F5496", "Blue", "ORDINANCES & STATUTES: ordinances, statutes, covenant signs, church practices, permissions, prologue"),
        ("1F4E79", "Navy", "TEACHING / PRINCIPLE: Christ's teaching on the law, prophetic and wisdom principles, the new-covenant promise"),
        ("7F3F00", "Brown", "JUDGMENT / CATALOGUE OF SINS: consequences pronounced (e.g. Genesis 3; Proverbs 6:16-19)"),
        ("7030A0", "Purple", "TRADITION / CUSTOM: human or rabbinic tradition, apostolic tradition, recorded custom"),
    ]:
        c = lg.cell(r[0], 1, label)
        c.font, c.border = Font(name=FONT, size=10, bold=True, color=color), BORDER
        lg.cell(r[0], 2, text).font = Font(name=FONT, size=9, color=color)
        r[0] += 1
    r[0] += 1
    head("3. Cell fill in the 'Classical division' column (Decalogue and Mosaic rows only)")
    for d, note in {
        "Moral": "Gold — abiding ethical law (Decalogue, worship, neighbour, family, oaths, idolatry)",
        "Ceremonial": "Blue — sanctuary, priesthood, sacrifice, purity, tithes, vows, signs",
        "Ceremonial — dietary": "Orange — clean/unclean food laws (a ceremonial sub-band, still debated today)",
        "Ceremonial — appointed times": "Light green — Sabbath, new moons, feasts (a ceremonial sub-band, still debated today)",
        "Civil / judicial": "Green — courts, crime, property, servants, kings and war, land, charity",
    }.items():
        c = lg.cell(r[0], 1, d)
        c.fill, c.font, c.border = PatternFill("solid", fgColor=DIV_FILL[d]), BOLD, BORDER
        lg.cell(r[0], 2, note).font = BODY
        r[0] += 1
    r[0] += 1
    c = lg.cell(r[0], 1, "Reading a row: shading tells you WHERE it comes from, font colour tells you WHAT KIND of rule it is, "
                         "and the division fill tells you HOW the classical threefold scheme classes it. The 'NT relation' and "
                         "'Status' columns are text tags — filter on them rather than reading colour.")
    c.font, c.alignment = SUB, Alignment(wrap_text=True, vertical="top")
    lg.merge_cells(start_row=r[0], start_column=1, end_row=r[0], end_column=2)
    lg.row_dimensions[r[0]].height = 30
    lg.column_dimensions["A"].width = 58
    lg.column_dimensions["B"].width = 100

    order = ["Summary", "Read Me", "Legend", "Master Matrix"] + [v[0] for v in VIEWS]
    wb._sheets = [wb[n] for n in order]
    wb.active = 0
    wb.save(out_path)
    print(f"{out_path}: {len(recs)} rows; " + "; ".join(f"{n} {c}" for n, c in view_counts))


def main():
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    version, repo_root = sys.argv[1], sys.argv[2]
    out_path = sys.argv[3] if len(sys.argv) > 3 else f"a_Bible_study_primer_commandments_laws_traditions_{version}.xlsx"
    build(version, repo_root, out_path)


if __name__ == "__main__":
    main()
