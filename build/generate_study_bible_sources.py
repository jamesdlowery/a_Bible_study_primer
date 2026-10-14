#!/usr/bin/env python3
"""Build the study-Bible sources workbook for a release:
a_Bible_study_primer_study_bible_sources_<version>.xlsx.

One row per study Bible that still needs a reader with the volume in hand
(the manual checklist's volumes): where to buy it (Amazon, Christianbook),
whether and where it exists electronically (Kindle, Logos, Olive Tree,
Accordance, the publisher's own app or ebook, any free PDF or web text),
and how many cells of that volume are still open on the manual checklist.
A volume whose count reaches zero is no longer needed and its row turns
green; the electronic-availability cell is shaded green / amber / red for
Full / Partial / None.

Sources:
  * build/verifications/study_bible_sources.csv -- the links, labels and
    notes (edit that file, never the spreadsheet). "Volume" must match the
    "## N. Name" heading in Section 120.
  * the live position tables and build/verifications/position_verifications.csv,
    via generate_manual_checklist.open_cells(), for the open-row counts --
    the same rows the manual checklist lists, so the two workbooks agree.

Every volume in the CSV must be a Section 120 heading, and every manual-
checklist volume should be in the CSV (a missing one is reported, with a
count, so it is not silently dropped). Usage:
    python3 build/generate_study_bible_sources.py <version> <repo_root> [out_path]
"""
import csv
import os
import sys
from collections import Counter

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import generate_manual_checklist as gm  # noqa: E402
import positions_matrix as pm  # noqa: E402

SOURCES_CSV = os.path.join("build", "verifications", "study_bible_sources.csv")

PLATFORMS = ["Kindle", "Logos", "Olive Tree", "Accordance", "Publisher app / ebook", "Free PDF / web"]
HEADERS = (["Study Bible", "Amazon", "Christianbook", "Open checklist rows", "Still needed?",
            "Electronic availability"] + PLATFORMS + ["Retail note", "Electronic note"])
WIDTHS = [36, 11, 14, 14, 13, 13, 18, 24, 18, 20, 28, 26, 55, 55]

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF")
PLAIN = Font(name=FONT)
INPUT = Font(name=FONT, color="0000FF")
LINK = Font(name=FONT, color="0563C1", underline="single")
MISSING = Font(name=FONT, italic=True, color="808080")
GREEN = PatternFill("solid", fgColor="C6EFCE")
AMBER = PatternFill("solid", fgColor="FFEB9C")
RED = PatternFill("solid", fgColor="FFC7CE")
_thin = Side(style="thin", color="BFBFBF")
BORDER = Border(top=_thin, bottom=_thin, left=_thin, right=_thin)


def read_sources(repo_root):
    with open(os.path.join(repo_root, SOURCES_CSV), encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def _link(cell, label, url):
    """A clickable cell: HYPERLINK when there is a URL, the bare label when
    there is only a label (e.g. 'NLT edition only'), 'none found' in grey
    when the CSV has neither."""
    label, url = (label or "").strip(), (url or "").strip()
    if url:
        cell.value = f'=HYPERLINK("{url}","{label or url}")'
        cell.font = LINK
    elif label:
        cell.value = label
        cell.font = PLAIN
    else:
        cell.value = "none found"
        cell.font = MISSING


def build(version, repo_root, out_path):
    rows = read_sources(repo_root)
    headings = set(pm.parse(repo_root)["Study Bibles"])
    unknown = [r["Volume"] for r in rows if r["Volume"] not in headings]
    if unknown:
        sys.exit(f"study_bible_sources.csv names volumes that are not Section 120 headings: {unknown}")
    counts = Counter(vol for vol, _cat, _pos, _src in gm.open_cells(repo_root))
    missing = sorted(set(counts) - {r["Volume"] for r in rows})
    for vol in missing:
        print(f"WARNING: {vol} has {counts[vol]} open checklist rows but no row in {SOURCES_CSV}")

    wb = Workbook()
    ws = wb.active
    ws.title = "Study Bibles"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for i, r in enumerate(rows, start=2):
        vol = r["Volume"]
        ws.cell(i, 1, vol).font = PLAIN
        _link(ws.cell(i, 2), "Amazon", r["Amazon URL"])
        _link(ws.cell(i, 3), "Christianbook", r["Christianbook URL"])
        ws.cell(i, 4, counts.get(vol, 0)).font = INPUT
        ws.cell(i, 5, f'=IF(D{i}=0,"No - complete","Yes")').font = PLAIN
        ws.cell(i, 6, r["Electronic availability"]).font = Font(name=FONT, bold=True)
        for j, p in enumerate(PLATFORMS, start=7):
            _link(ws.cell(i, j), r[p], r[p + " URL"])
        ws.cell(i, 13, r["Retail note"]).font = PLAIN
        ws.cell(i, 14, r["Electronic note"]).font = PLAIN
        for col in range(1, len(HEADERS) + 1):
            c = ws.cell(i, col)
            c.border = BORDER
            c.alignment = Alignment(wrap_text=True, vertical="top",
                                    horizontal="center" if col in (4, 5, 6) else "general")

    last = len(rows) + 1
    ws.conditional_formatting.add(f"A2:N{last}", FormulaRule(formula=["$D2=0"], fill=GREEN))
    ws.conditional_formatting.add(f"F2:F{last}", FormulaRule(formula=['$F2="Full"'], fill=GREEN))
    ws.conditional_formatting.add(f"F2:F{last}", FormulaRule(formula=['$F2="Partial"'], fill=AMBER))
    ws.conditional_formatting.add(f"F2:F{last}", FormulaRule(formula=['$F2="None"'], fill=RED))
    dv = DataValidation(type="list", formula1='"Full,Partial,None"')
    ws.add_data_validation(dv)
    dv.add(f"F2:F{last}")
    for i, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:N{last}"

    s = last + 2
    ws.cell(s, 1, "Summary").font = Font(name=FONT, bold=True)
    summary = [("Volumes still needed", f'=COUNTIF($D$2:$D${last},">0")'),
               ("Volumes with verification complete", f"=COUNTIF($D$2:$D${last},0)"),
               ("Open checklist rows in all", f"=SUM($D$2:$D${last})"),
               ("Full electronic edition", f'=COUNTIF($F$2:$F${last},"Full")'),
               ("Partial electronic edition", f'=COUNTIF($F$2:$F${last},"Partial")'),
               ("No electronic edition found", f'=COUNTIF($F$2:$F${last},"None")')]
    for k, (label, formula) in enumerate(summary, start=1):
        ws.cell(s + k, 1, label).font = PLAIN
        c = ws.cell(s + k, 2, formula)
        c.font, c.alignment = PLAIN, Alignment(horizontal="center")

    n = s + len(summary) + 2
    ws.cell(n, 1, "How to read").font = Font(name=FONT, bold=True)
    notes = [
        f"Generated with build {version} from build/verifications/study_bible_sources.csv (links, labels, notes "
        f"-- edit that file, not this workbook) and the live position tables with "
        f"build/verifications/position_verifications.csv (the open-row counts).",
        "Column D is the number of rows for that volume on the 'Study Bibles - notes to check' sheet of the "
        "manual checklist workbook built in the same run. At 0 the whole row turns green and column E reads "
        "'No - complete': the volume is no longer needed.",
        "Column F: Full = a complete study Bible (text + notes) or the full notes module on at least one "
        "platform; Partial = notes only for an older edition, part of the notes, or restricted library access; "
        "None = no legitimate electronic edition found. Only the F cell is shaded green / amber / red.",
        "Platform cells: 'Full' = text and notes together; 'Notes only' = notes module sold separately from "
        "the Bible text; 'none found' = nothing surfaced when the CSV row was researched, not proof that no "
        "edition exists. Kindle links were taken from search results and should be opened to confirm the "
        "edition. 'Logos' the platform (Faithlife) is unrelated to Logos International, publisher of the 1972 "
        "Logos Complete Study Bible.",
    ]
    for j, t in enumerate(notes, start=1):
        c = ws.cell(n + j, 1, t)
        c.font, c.alignment = PLAIN, Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=n + j, start_column=1, end_row=n + j, end_column=len(HEADERS))
        ws.row_dimensions[n + j].height = 30

    wb.save(out_path)
    done = sum(1 for r in rows if counts.get(r["Volume"], 0) == 0)
    print(f"Wrote {out_path}: {len(rows)} volumes, {done} complete, "
          f"{sum(counts.values())} open checklist rows")


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    version = sys.argv[1]
    repo_root = sys.argv[2] if len(sys.argv) > 2 else "."
    out_path = sys.argv[3] if len(sys.argv) > 3 else f"a_Bible_study_primer_study_bible_sources_{version}.xlsx"
    build(version, repo_root, out_path)


if __name__ == "__main__":
    main()
